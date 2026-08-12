#!/usr/bin/env python3
"""Generate complete OBPS module technical Word document (routes, roles, files, functions)."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "OBPS_Module_Complete_Technical_Document.docx"
IMG = ROOT / "_lld_assets"
IMG.mkdir(exist_ok=True)

PRIMARY = (11, 61, 92)
ACCENT = (31, 111, 139)

# ---------------------------------------------------------------------------
# File task catalog (every source file)
# ---------------------------------------------------------------------------
TASKS = {
    "src/Module.js": (
        "Citizen + Employee + Architect",
        "Module bootstrap. Loads i18n store for bpa/bpareg/common, stores OBPS_TENANTS, "
        "renders CitizenApp or EmployeeApp by userType. Registers all components via initOBPSComponents. "
        "OBPSLinks shows citizen home cards: Stakeholder registration vs Architect login.",
    ),
    "src/pages/citizen/index.js": (
        "Citizen + Architect",
        "Citizen router. Declares every citizen OBPS route under /upyog-ui/citizen/obps. "
        "Wraps most screens in PrivateRoute. Hides BackButton on acknowledgement/response and after EDCR apply.",
    ),
    "src/pages/employee/index.js": (
        "Employee",
        "Employee router + breadcrumbs. Routes inbox, stakeholder inbox, search, BPA/stakeholder detail, "
        "workflow responses, and EnhancedReport (daily collection + application status).",
    ),
    "src/pages/citizen/home.js": (
        "Architect (stakeholder on citizen portal)",
        "Architect home. MDMS StakeholderRegistraition.TradeTypetoRoleMapping gates access; "
        "non-stakeholder users get BPA_LOGIN_HOME_VALIDATION_MESSAGE_LABEL toast and redirect to all-services. "
        "Shows BPA inbox KPI, EDCR inbox KPI, eDCR apply links, and MDMS homePageUrlLinks for BPA/OC apply.",
    ),
    "src/pages/employee/EmployeeCard.js": (
        "Employee",
        "Employee home module card. Shows Stakeholder inbox if BPAREG_DOC_VERIFIER/BPAREG_APPROVER; "
        "BPA inbox if BPA_FIELD_INSPECTOR/BPA_NOC_VERIFIER/BPA_APPROVER/BPA_VERIFIER/CEMP. "
        "KPIs: total inbox count and nearing SLA.",
    ),
    "src/pages/citizen/EDCR/index.js": (
        "Architect",
        "New-construction eDCR wizard. Session EDCR_CREATE. handleSelect builds multipart edcrRequest "
        "(BUILDING_PLAN_SCRUTINY / NEW_CONSTRUCTION) and calls Digit.EDCRService.create.",
    ),
    "src/pages/citizen/EDCR/EDCRAcknowledgement.js": (
        "Architect",
        "eDCR result: ERROR / Accepted / Rejected. Download planReport; Accepted CTA Apply for BPA via MDMS homePageUrlLinks.",
    ),
    "src/pages/citizen/OCEDCR/index.js": (
        "Architect",
        "OC eDCR wizard. createOCEdcr posts BUILDING_OC_PLAN_SCRUTINY / NEW_CONSTRUCTION with plan file.",
    ),
    "src/pages/citizen/OCEDCR/EDCRAcknowledgement.js": (
        "Architect",
        "OC eDCR ack. Accepted CTA Apply OC BPA. Download OC scrutiny report.",
    ),
    "src/pages/citizen/NewBuildingPermit/index.js": (
        "Architect",
        "BPA apply wizard. Uses NewConfig.js steps (includes CPT property search). Session BUILDING_PERMIT. "
        "Wildcard Navigate to docs-required. After last step → check → acknowledgement.",
    ),
    "src/pages/citizen/NewBuildingPermit/NewConfig.js": (
        "Architect",
        "Active BPA step list: docs-required → … → CPT property → location → owners → documents → noc → check.",
    ),
    "src/pages/citizen/NewBuildingPermit/CheckPage.js": (
        "Architect",
        "BPA review/summary before submit. Shows scrutiny, plot, owners, docs, NOC, fees. Submit → acknowledgement mutation.",
    ),
    "src/pages/citizen/NewBuildingPermit/OBPSAcknowledgement.js": (
        "Architect",
        "BPA create/update ack. Downloads PDF via getBPAAcknowledgement. Messages keyed by businessService/action/status/architect type.",
    ),
    "src/pages/citizen/OCBuildingPermit/index.js": (
        "Architect",
        "OC BPA wizard from ocbuildingPermitConfig. Session BUILDING_PERMIT with OC eDCR number.",
    ),
    "src/pages/citizen/OCBuildingPermit/CheckPage.js": (
        "Architect",
        "OC BPA review/summary (building extract, demolition, docs, NOC).",
    ),
    "src/pages/citizen/OCBuildingPermit/OBPSAcknowledgement.js": (
        "Architect",
        "OC BPA acknowledgement banner and PDF download.",
    ),
    "src/pages/citizen/StakeholderRegistration/index.js": (
        "Citizen (also open-link unauthenticated)",
        "Stakeholder (architect/engineer/…) registration wizard from stakeholderConfig. "
        "Logged-in: /stakeholder/apply/* ; public: /openlink/stakeholder/apply/*.",
    ),
    "src/pages/citizen/StakeholderRegistration/CheckPage.js": (
        "Citizen",
        "Stakeholder registration review: license type, personal details, addresses, documents, fee estimate.",
    ),
    "src/pages/citizen/StakeholderRegistration/StakeholderAcknowledgement.js": (
        "Citizen",
        "BPAREG create ack. Proceed next / download acknowledgement PDF (getAcknowlegment.js).",
    ),
    "src/pages/citizen/PreApprovedPlan/index.js": (
        "Architect",
        "Pre-approved plan (PAP) wizard. Session BPA_PRE_APPROVED_CREATE. Starts at documents-required.",
    ),
    "src/pages/citizen/MyApplication/index.js": (
        "Citizen + Architect",
        "My applications list: BPAREG licenses + BPA applications. Architects can resume INITIATED via getBPAFormData.",
    ),
    "src/pages/citizen/ApplicationDetail/index.js": (
        "Citizen + Architect",
        "Stakeholder registration application detail (BPAREG) for citizen: license, owner, docs, timeline.",
    ),
    "src/pages/citizen/BpaApplicationDetail/index.js": (
        "Citizen + Architect",
        "Citizen BPA application detail. Workflow actions, payments, permit/OC PDFs, T&C checkbox. "
        "Architect action bar hidden during CITIZEN_APPROVAL_INPROCESS. Citizen can SEND_TO_ARCHITECT or approve.",
    ),
    "src/pages/citizen/BpaApplicationDetail/BPAApplicationTimeline.js": (
        "Citizen + Architect",
        "Workflow timeline captions including architect/citizen checkpoints and image zoom.",
    ),
    "src/pages/citizen/BpaApplicationDetail/BPACaption.js": (
        "Citizen + Architect",
        "Single timeline caption row (assignee, comments, documents).",
    ),
    "src/pages/citizen/BpaApplicationDetail/BPAReason.js": (
        "Citizen + Architect",
        "Displays headComment/otherComment for rejection or send-back reasons.",
    ),
    "src/pages/citizen/BpaApplicationDetail/Modal/index.js": (
        "Citizen + Architect",
        "Action modal: comments, optional file, submit workflow action (FORWARD/SEND_BACK/APPROVE/etc.).",
    ),
    "src/pages/citizen/BpaApplicationDetail/config/Approve.js": (
        "Citizen + Architect",
        "Form config for accept/forward application (comment + upload).",
    ),
    "src/pages/citizen/BpaApplicationDetail/config/TermsAndConditions.js": (
        "Citizen",
        "T&C config shown before citizen approval of BPA.",
    ),
    "src/pages/citizen/BPASendToArchitect/index.js": (
        "Citizen",
        "Citizen edits application then SEND_TO_ARCHITECT. Loads existing BPA via getBPAEditDetails into wizard.",
    ),
    "src/pages/citizen/OCSendToArchitect/index.js": (
        "Citizen",
        "OC equivalent of send-to-architect edit wizard.",
    ),
    "src/pages/citizen/BPASendBackToCitizen/index.js": (
        "Architect",
        "Architect resumes application after SEND_BACK_TO_CITIZEN / citizen resubmit path.",
    ),
    "src/pages/citizen/OCSendBackToCitizen/index.js": (
        "Architect",
        "OC send-back-to-citizen resume wizard.",
    ),
    "src/pages/citizen/OCSendBackToCitizen/CheckPage.js": (
        "Architect",
        "OC send-back review with workflow FORWARD action from check page.",
    ),
    "src/pages/citizen/OCSendBackToCitizen/Acknowledgement.js": (
        "Architect",
        "OC send-back acknowledgement messages using BPA_SUBMIT_APP session flag.",
    ),
    "src/pages/citizen/ArchitectInbox/index.js": (
        "Architect",
        "Architect BPA inbox (also used historically). Searches bpa-services BPA_LOW/BPA/BPA_OC. Pagination/filter/sort.",
    ),
    "src/pages/citizen/ArchitectInbox/DesktopInbox.js": (
        "Architect",
        "Desktop table for architect inbox (app no, type, service, SLA, status).",
    ),
    "src/pages/citizen/ArchitectInbox/MobileInbox.js": (
        "Architect",
        "Mobile cards combining BPA + EDCR + BPAREG rows.",
    ),
    "src/pages/citizen/ArchitectInbox/ApplicationCard.js": (
        "Architect",
        "Mobile card wrapper with search/filter/sort popups.",
    ),
    "src/pages/citizen/ArchitectInbox/ApplicationTable.js": (
        "Architect",
        "Reusable table pagination for architect inbox.",
    ),
    "src/pages/citizen/ArchitectInbox/ApplicationLinks.js": (
        "Architect",
        "Quick links header on architect inbox.",
    ),
    "src/pages/citizen/ArchitectInbox/Filter.js": (
        "Architect",
        "Filter by status, application type, risk type.",
    ),
    "src/pages/citizen/ArchitectInbox/Search.js": (
        "Architect",
        "Search form fields for architect inbox.",
    ),
    "src/pages/citizen/ArchitectInbox/SortBy.js": (
        "Architect",
        "Sort options for architect inbox.",
    ),
    "src/pages/citizen/ArchitectInbox/Status.js": (
        "Architect",
        "Status checklist in filter.",
    ),
    "src/pages/citizen/ArchitectInbox/StatusCount.js": (
        "Architect",
        "Per-status count checkbox row.",
    ),
    "src/pages/citizen/EdcrInbox/index.js": (
        "Architect",
        "eDCR inbox. Links to apply and OC apply. Session EDCR.INBOX. useEDCRInbox hook.",
    ),
    "src/pages/citizen/EdcrInbox/FilterFormFieldsComponent.js": (
        "Architect",
        "eDCR inbox filters (status, service type).",
    ),
    "src/pages/citizen/EdcrInbox/SearchFormFieldsComponent.js": (
        "Architect",
        "eDCR search by application number / eDCR number.",
    ),
    "src/pages/citizen/EdcrInbox/useInboxTableConfig.js": (
        "Architect",
        "Desktop columns: app no, eDCR no, status, download DXF/report.",
    ),
    "src/pages/citizen/EdcrInbox/useInboxMobileCardsData.js": (
        "Architect",
        "Mobile cards for eDCR inbox rows.",
    ),
    "src/pages/employee/Inbox/index.js": (
        "Employee (BPA roles)",
        "Employee BPA inbox. moduleName bpa-services. Session OBPS.INBOX. Search by mobile/app no; filter status/type/assignee.",
    ),
    "src/pages/employee/Inbox/FilterFormFieldsComponent.js": (
        "Employee",
        "Inbox filters including application type checkboxes and locality.",
    ),
    "src/pages/employee/Inbox/SearchFormFieldsComponent.js": (
        "Employee",
        "Search application number and mobile number.",
    ),
    "src/pages/employee/Inbox/useInboxTableConfig.js": (
        "Employee",
        "Employee inbox table columns and row link to /inbox/bpa/:id.",
    ),
    "src/pages/employee/Inbox/useInboxMobileCardsData.js": (
        "Employee",
        "Mobile cards for employee BPA inbox.",
    ),
    "src/pages/employee/stakeholderInbox/index.js": (
        "Employee (BPAREG roles)",
        "Stakeholder registration inbox (BPAREG). Session STAKEHOLDER.INBOX. Roles BPAREG_EMPLOYEE/APPROVER/DOC_VERIFIER.",
    ),
    "src/pages/employee/stakeholderInbox/FilterFormFieldsComponent.js": (
        "Employee",
        "BPAREG inbox filters including license type.",
    ),
    "src/pages/employee/stakeholderInbox/SearchFormFieldsComponent.js": (
        "Employee",
        "BPAREG search by application number.",
    ),
    "src/pages/employee/stakeholderInbox/useInboxTableConfig.js": (
        "Employee",
        "Stakeholder inbox table → /stakeholder-inbox/stakeholder/:id.",
    ),
    "src/pages/employee/stakeholderInbox/useInboxMobileCardsData.js": (
        "Employee",
        "Mobile cards for stakeholder inbox.",
    ),
    "src/pages/employee/Search.js": (
        "Employee + Architect (citizen search routes reuse this)",
        "Application search page wrapping OBPSSearchApplication. Citizen vs employee query params differ.",
    ),
    "src/pages/employee/ApplicationDetail/index.js": (
        "Employee",
        "Employee stakeholder (BPAREG) application detail using ApplicationDetails template + workflow actions.",
    ),
    "src/pages/employee/BpaApplicationDetails/index.js": (
        "Employee",
        "Employee BPA/OC/PAP detail. Downloads receipts, permit/OC order, revocation PDF, comparison report. "
        "Field inspector sees InspectionReport form (hideInCitizen). Workflow via ApplicationDetailsTemplate.",
    ),
    "src/pages/employee/OBPSResponse.js": (
        "Employee (+ citizen /response)",
        "Post-workflow response. Fetches BPA + bill; Pay CTA for citizen/employee; revocation PDF if applicable.",
    ),
    "src/pages/employee/StakeholderResponse.js": (
        "Employee",
        "Post BPAREG workflow response header/subheader and go home.",
    ),
    "src/pageComponents/EDCRForm.js": (
        "Architect",
        "eDCR apply/home form: city (BPAAPPLY tenants), applicant name, plan file upload.",
    ),
    "src/pageComponents/DocsRequired.js": (
        "Architect",
        "Info screen listing required documents from MDMS DocumentTypes. Sets BPA_ARCHITECT_NAME from approved license.",
    ),
    "src/pageComponents/PreApprovedDocsRequired.js": (
        "Architect",
        "PAP required-documents info (plot dimensions, submission requirements).",
    ),
    "src/pageComponents/OCEDCRDocsRequired.js": (
        "Architect",
        "OC eDCR required inputs info (permit no, DXF).",
    ),
    "src/pageComponents/StakeholderDocsRequired.js": (
        "Citizen",
        "Stakeholder registration required documents from MDMS.",
    ),
    "src/pageComponents/BasicDetails.js": (
        "Architect",
        "eDCR number search via scrutinyDetailsData; shows application type, occupancy, risk type, applicant. "
        "8-char drawing no uses PreApprovedPlanService.",
    ),
    "src/pageComponents/OCBasicDetails.js": (
        "Architect",
        "OC basic details: permit city validation, OCeDCR search, risk-type mismatch error/warning vs permit.",
    ),
    "src/pageComponents/PlotDetails.js": (
        "Architect",
        "Plot/katha/holding number and land registration details.",
    ),
    "src/pageComponents/ScrutinyDetails.js": (
        "Architect",
        "Plan scrutiny display: blocks, floors, sub-occupancy multi-select, demolition area, building extract. "
        "Can call Digit.OBPSService.create for INITIATED draft.",
    ),
    "src/pageComponents/LocationDetails.js": (
        "Architect",
        "Address: city, locality, pincode, street, landmark, GIS pin (GIS component).",
    ),
    "src/pageComponents/GIS.js": (
        "Architect",
        "Map pin picker; validates pincode serviceability.",
    ),
    "src/pageComponents/OwnerDetails.js": (
        "Architect",
        "Owners: add/remove, gender, mobile, email, primary owner. User search/create. "
        "On last owner step may Digit.OBPSService.create INITIATED BPA. Blocks CITIZEN-only users from architect flows.",
    ),
    "src/pageComponents/DocumentDetails.js": (
        "Architect",
        "Upload BPA documents per MDMS DocTypeMapping.",
    ),
    "src/pageComponents/NOCDetails.js": (
        "Architect",
        "NOC documents per NOC type; SelectDocument upload helper.",
    ),
    "src/pageComponents/BuildingPlanScrutiny.js": (
        "Architect",
        "PAP drawing search, thumbnails, estimate/fee rows, CAD/PDF/image diagram preview, proceed.",
    ),
    "src/pageComponents/OCeDCRScrutiny.js": (
        "Architect",
        "Search permit number + date; load BPA + eDCR; proceed to OC diagram upload.",
    ),
    "src/pageComponents/OCUploadPlanDiagram.js": (
        "Architect",
        "Upload OC plan DXF; submit to parent createOCEdcr.",
    ),
    "src/pageComponents/LicenseType.js": (
        "Citizen",
        "Select trade type from TradeTypetoRoleMapping (Architect/Engineer/…). Council number mandatory for ARCHITECT.",
    ),
    "src/pageComponents/LicenseDetails.js": (
        "Citizen",
        "Stakeholder personal details: name, gender, mobile, email, PAN.",
    ),
    "src/pageComponents/PermanentAddress.js": (
        "Citizen",
        "Permanent address text for licensee.",
    ),
    "src/pageComponents/CorrospondenceAddress.js": (
        "Citizen",
        "Correspondence address; optional same-as-permanent. May BPAREGCreate draft.",
    ),
    "src/pageComponents/StakeholderDocuments.js": (
        "Citizen",
        "Upload stakeholder registration documents.",
    ),
    "src/pageComponents/InspectionReport.js": (
        "Employee (Field Inspector)",
        "Field inspection reports: date, checklist, remarks, documents. hideInCitizen. Multiple reports via addNewFieldReport.",
    ),
    "src/pageComponents/OBPSDocuments.js": (
        "Citizen + Architect + Employee",
        "Read-only document preview by Code (BPA/NOC/stakeholder).",
    ),
    "src/pageComponents/OBPSDocumentsEmp.js": (
        "Employee",
        "Employee document upload during scrutiny/inspection.",
    ),
    "src/pageComponents/OBPSDocumentsHolder.js": (
        "Citizen + Architect + Employee",
        "Groups documents into titled Document lists.",
    ),
    "src/components/Timeline.js": (
        "Citizen + Architect",
        "Stepper UI. Flows: default BPA, STAKEHOLDER, OCBPA, PRE_APPROVE.",
    ),
    "src/components/ApplicationTimeline.js": (
        "Citizen + Architect",
        "Generic workflow timeline with captions and next actions.",
    ),
    "src/components/DocumentDetails.js": (
        "Citizen + Architect + Employee",
        "Simple document list renderer.",
    ),
    "src/components/SearchApplication/index.js": (
        "Employee + Architect",
        "Full search UI: form, table, mobile modal, pagination, redirect by businessService.",
    ),
    "src/components/SearchApplication/SearchFormFieldsComponent.js": (
        "Employee + Architect",
        "Search fields: app no, mobile, dates, type, service, status. Role-aware BPA vs BPAREG.",
    ),
    "src/components/SearchApplication/useTableConfig.js": (
        "Employee + Architect",
        "Search results table column definitions.",
    ),
    "src/config/edcrConfig.js": (
        "Architect",
        "Default eDCR step: home → EDCRForm.",
    ),
    "src/config/ocEdcrConfig.js": (
        "Architect",
        "OC eDCR steps: docs-required → home → upload-diagram.",
    ),
    "src/config/buildingPermitConfig.js": (
        "Architect",
        "Fallback BPA steps (without CPT). Overridden by MDMS BuildingPermitConfig; runtime uses NewConfig.js.",
    ),
    "src/config/ocbuildingPermitConfig.js": (
        "Architect",
        "OC BPA steps: docs-required → OCBasicDetails → plot → scrutiny → documents → noc.",
    ),
    "src/config/stakeholderConfig.js": (
        "Citizen",
        "BPAREG steps: docs → license type → details → permanent → correspondence → documents.",
    ),
    "src/config/PreApprovedPlanConfig.js": (
        "Architect",
        "PAP steps: documents-required → planDetails → BasicDetails → plot → scrutiny → location → owners → documents.",
    ),
    "src/config/InspectionReportConfig.js": (
        "Employee",
        "Employee-only InspectionReport component on BPA detail.",
    ),
    "src/utils/index.js": (
        "All roles",
        "Shared helpers: validation patterns, BPA/OC/stakeholder payload converters, scrutiny lookups, "
        "PDF download/print, role-link visibility, fee businessService mapping.",
    ),
    "getAcknowlegment.js": (
        "Citizen",
        "Builds PDF data payload for stakeholder registration acknowledgement.",
    ),
    "getBPAAcknowledgement.js": (
        "Architect + Citizen",
        "Builds PDF data payload for BPA/OC/PAP acknowledgement (owners, plot, location, scrutiny).",
    ),
}

UTILS_FUNCS = {
    "getPattern": "Regex for Name (no digits/specials, max 50) or MobileNo (6-9XXXXXXXXX).",
    "stringReplaceAll": "Replace all occurrences of a substring.",
    "sortDropdownNames": "Sort options by localized label.",
    "uuidv4": "Generate UUID via uuid package (eDCR transactionNumber).",
    "pdfDownloadLink": "Pick original (non thumbnail) file URL from filestore map.",
    "convertToNocObject": "Build Noc payload with documents filtered by nocType.",
    "getBPAFormData": "Load eDCR + map BPA into wizard session; navigate to bpa or ocbpa path.",
    "getDocumentforBPA": "Merge current + previous-state documents for API.",
    "getusageCategoryAPI": "Join occupancy codes with commas.",
    "getBPAUnit": "Merge landInfo.unit with subOccupancy selections.",
    "getBPAusageCategoryAPI": "Same as getusageCategoryAPI (exported).",
    "getBPAUnitsForAPI": "Build unit array from subOccupancy object keys Block_N.",
    "getunitforBPA": "Normalize existing units for API.",
    "getBPAOwners": "Merge owners; OC/INITIATED returns landInfo.owners; deactivate removed owners.",
    "getOwnerShipCategory": "Ownership category code from form or landInfo.",
    "convertToBPAObject": "Full BPA create/update payload; workflow APPLY for BPA-PAP else SEND_TO_CITIZEN.",
    "getapplicationdocstakeholder": "Map stakeholder docs to applicationDocuments.",
    "convertToStakeholderObject": "Licenses APPLY payload with tradeType, council no, owners, addresses.",
    "getUniqueItemsFromArray": "Dedupe by identifier.",
    "convertDateToEpoch": "YYYY-MM-DD to epoch (dayend optional).",
    "convertEpochToDateDMY": "Epoch to DD/MM/YYYY.",
    "getBPAEditDetails": "Hydrate edit-application wizard from existing BPA + eDCR + NOC MDMS.",
    "getPath": "Replace :param tokens in a path.",
    "convertDateTimeToEpoch": "Datetime string to epoch.",
    "convertEpochToDate": "Epoch to YYYY-MM-DD.",
    "getBusinessServices": "Map BPA_LOW/BPA/BPA-PAP/BPA_OC + status to bill businessService (APP_FEE vs SAN_FEE).",
    "downloadPdf": "Download blob (mSewa or anchor).",
    "printPdf": "Open blob URL and print.",
    "downloadAndPrintReciept": "Generate or reuse payment PDF and open.",
    "getOrderedDocs": "Group docs by documentType with filestoreIdArray.",
    "showHidingLinksForStakeholder": "True if employee has BPAREG role+tenant match.",
    "showHidingLinksForBPA": "True if employee has BPA/CEMP role on current tenant.",
    "getCheckBoxLabelData": "T&C label: citizen vs stakeholder declaration by status.",
    "scrutinyDetailsData": "Lookup eDCR or PAP drawing; error if BPA already exists (unless INITIATED/REJECTED/PERMIT REVOCATION).",
    "getOCEDCRDetails": "OC eDCR scrutinyDetails wrapper.",
    "ocScrutinyDetailsData": "OC eDCR + linked permit BPA + original eDCR bundle.",
    "getOrderDocuments": "Group uploaded docs by type prefix for preview.",
}


def font(size, bold=False):
    for path in (
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ):
        try:
            return ImageFont.truetype(path, size)
        except Exception:
            continue
    return ImageFont.load_default()


def make_role_diagram():
    w, h = 1400, 780
    img = Image.new("RGB", (w, h), (248, 250, 252))
    d = ImageDraw.Draw(img)
    d.text((40, 20), "OBPS Role × Portal × Primary Tasks", font=font(22, True), fill=PRIMARY)

    boxes = [
        (40, 80, 450, 360, (11, 61, 92), "CITIZEN",
         ["Portal: /upyog-ui/citizen", "Stakeholder self-registration", "  /stakeholder/apply", "  /openlink/stakeholder/apply",
          "My applications", "Approve BPA (T&C)", "Send application to Architect", "Pay application / sanction fee"]),
        (475, 80, 925, 360, (31, 111, 139), "ARCHITECT / STAKEHOLDER",
         ["Same citizen portal after BPAREG", "Roles: BPA_ARCHITECT, ENGINEER,", "  BUILDER, STRUCTURALENGINEER,", "  SUPERVISOR, TOWNPLANNER",
          "eDCR + OC eDCR scrutiny", "Apply BPA / OCBPA / PAP", "Architect inbox + EDCR inbox", "Send back / resume applications"]),
        (950, 80, 1360, 360, (120, 50, 40), "EMPLOYEE",
         ["Portal: /upyog-ui/employee", "BPA: VERIFIER, APPROVER,", "  FIELD_INSPECTOR, NOC_VERIFIER, CEMP",
          "BPAREG: DOC_VERIFIER, APPROVER", "Inbox / search / reports", "Inspection report", "Permit / OC / receipts"]),
    ]
    for x0, y0, x1, y1, color, title, lines in boxes:
        d.rounded_rectangle((x0, y0, x1, y1), 12, fill=(255, 255, 255), outline=color, width=3)
        d.rectangle((x0 + 2, y0 + 2, x1 - 2, y0 + 42), fill=color)
        tb = d.textbbox((0, 0), title, font=font(14, True))
        d.text((x0 + (x1 - x0 - (tb[2] - tb[0])) / 2, y0 + 12), title, font=font(14, True), fill=(255, 255, 255))
        y = y0 + 56
        for line in lines:
            d.text((x0 + 16, y), line, font=font(13), fill=(30, 30, 30))
            y += 22

    d.rounded_rectangle((40, 400, 1360, 740), 12, fill=(255, 255, 255), outline=PRIMARY, width=2)
    d.text((60, 420), "End-to-end (happy path)", font=font(16, True), fill=PRIMARY)
    flow = [
        "1. Citizen registers as Architect (BPAREG) → Employee verifies/approves license",
        "2. Architect logs into citizen OBPS home → eDCR apply/home → scrutinize DXF → Accepted eDCR number",
        "3. Apply BPA (docs, scrutiny, property, location, owners, docs, NOC) → SEND_TO_CITIZEN",
        "4. Citizen reviews T&C and approves / pays app fee  OR  sends back to architect",
        "5. Employee verifier / field inspector / NOC / approver process inbox → permit order",
        "6. Later: OC eDCR + OC BPA for occupancy certificate",
    ]
    y = 460
    for line in flow:
        d.text((60, y), line, font=font(14), fill=(30, 30, 30))
        y += 40
    path = IMG / "obps_roles.png"
    img.save(path)
    return path


def extract_files():
    files = sorted(p for p in ROOT.rglob("*") if p.suffix in {".js", ".jsx"} and "node_modules" not in str(p) and "_lld_assets" not in str(p) and p.name.startswith("generate_") is False)
    out = []
    for p in files:
        rel = str(p.relative_to(ROOT)).replace("\\", "/")
        text = p.read_text(errors="replace")
        fns = []
        for m in re.finditer(r"(?:export\s+)?(?:async\s+)?function\s+(\w+)\s*\(([^)]*)\)", text):
            fns.append((m.group(1), re.sub(r"\s+", " ", m.group(2).strip())[:140]))
        for m in re.finditer(r"(?:export\s+)?const\s+(\w+)\s*=\s*(?:async\s*)?\(([^)]*)\)\s*=>", text):
            fns.append((m.group(1), re.sub(r"\s+", " ", m.group(2).strip())[:140]))
        for m in re.finditer(r"(?:export\s+)?const\s+(\w+)\s*=\s*(?:async\s+)?function\s*\(([^)]*)\)", text):
            fns.append((m.group(1), re.sub(r"\s+", " ", m.group(2).strip())[:140]))
        # unique preserve order
        seen = set()
        uniq = []
        for n, a in fns:
            if n not in seen:
                seen.add(n)
                uniq.append((n, a))
        apis = sorted(set(re.findall(r"Digit\.(?:[A-Za-z]+Service)\.[A-Za-z0-9_]+", text)))
        hooks = sorted(set(re.findall(r"Digit\.Hooks(?:\.[A-Za-z0-9_]+)+", text)))[:15]
        out.append({
            "rel": rel,
            "lines": text.count("\n") + 1,
            "fns": uniq,
            "apis": apis[:20],
            "hooks": hooks,
        })
    return out


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=PRIMARY)


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=11)


def shade(cell, fill="0B3D5C"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    cell._element.get_or_add_tcPr().append(shd)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=9, bold=True, color=(255, 255, 255))
        shade(cell)
    for r_idx, row in enumerate(rows):
        fill = "F3F7FA" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val) if val is not None else "")
            set_run_font(run, size=8)
            shade(cell, fill)
    doc.add_paragraph()
    return table


def infer_role(rel, catalog_role):
    if catalog_role:
        return catalog_role
    if "employee" in rel.lower() or "Emp" in rel or "Inspection" in rel:
        return "Employee"
    if "Stakeholder" in rel or "License" in rel or "Permanent" in rel or "Corrospondence" in rel:
        return "Citizen"
    if "citizen" in rel or "pageComponents" in rel or "EDCR" in rel:
        return "Architect"
    return "All roles"


def main():
    files = extract_files()
    role_img = make_role_diagram()

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.65)
        s.bottom_margin = Inches(0.65)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("UPYOG OBPS Module")
    set_run_font(r, 26, True, PRIMARY)
    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("Complete Technical Document\nRoutes • Roles • Files • Functions")
    set_run_font(r, 16, True, ACCENT)
    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = m.add_run(
        "Package: @upyog/digit-ui-module-obps  v3.10.0\n"
        "Path: frontend/upyog-ui/web/micro-ui-internals/packages/modules/obps\n"
        "Entry: src/Module.js  |  Citizen base: /upyog-ui/citizen/obps  |  Employee base: /upyog-ui/employee/obps\n"
        f"Source files documented: {len(files)}"
    )
    set_run_font(r, 10, False, (80, 80, 80))

    add_heading(doc, "1. Purpose and scope", 1)
    add_para(
        doc,
        "This document inventories the entire OBPS (Online Building Plan Scrutiny / Building Permit) UI module. "
        "It covers citizen, architect/stakeholder, and employee experiences: every route, wizard step, source file, "
        "exported/inner function, API call, session key, MDMS master, and role gate found in this package. "
        "Nothing in the module source tree (excluding node_modules and generator scripts) is omitted from Section 8.",
    )
    add_bullet(doc, "Citizen: unlicensed public user — primarily stakeholder (BPAREG) self-registration, later BPA approval and payment.")
    add_bullet(doc, "Architect: licensed stakeholder on the citizen portal (BPA_ARCHITECT and sibling trade roles) — eDCR, BPA, OC, PAP, inboxes.")
    add_bullet(doc, "Employee: ULB staff on employee portal — BPAREG verification, BPA scrutiny/inspection/approval, reports.")

    add_heading(doc, "2. Module architecture", 1)
    add_para(doc, "initOBPSComponents() registers every screen/component on Digit.ComponentRegistryService. OBPSModule chooses CitizenApp vs EmployeeApp from userType.")
    add_table(
        doc,
        ["Layer", "Location", "Responsibility"],
        [
            ("Bootstrap", "src/Module.js", "i18n store, tenants, registry, Citizen vs Employee split"),
            ("Citizen routes", "src/pages/citizen/index.js", "All /citizen/obps routes"),
            ("Employee routes", "src/pages/employee/index.js", "All /employee/obps routes"),
            ("Wizards", "pages/citizen/* index.js", "Multi-step apply flows + session"),
            ("Steps", "src/pageComponents/", "Individual form screens used by wizards"),
            ("Configs", "src/config/ + NewConfig.js", "Step order; MDMS can override"),
            ("Shared UI", "src/components/", "Timeline, search, document preview"),
            ("Helpers", "src/utils/index.js", "Payload mapping, PDF, role checks, scrutiny"),
            ("Ack PDF", "getAcknowlegment.js, getBPAAcknowledgement.js", "PDF JSON for receipts"),
        ],
    )

    add_heading(doc, "3. Roles in detail", 1)
    doc.add_picture(str(role_img), width=Inches(6.6))
    doc.add_paragraph()

    add_heading(doc, "3.1 Citizen", 2)
    add_para(doc, "Portal user without a BPAREG professional license (or before approval). Typical role code on user object: CITIZEN.")
    add_bullet(doc, "Register as stakeholder (Architect/Engineer/Builder/Structural Engineer/Supervisor/Town Planner) via TradeTypetoRoleMapping.")
    add_bullet(doc, "Open-link flow /openlink/stakeholder/apply/* does not require PrivateRoute (public).")
    add_bullet(doc, "After BPA is SEND_TO_CITIZEN: status CITIZEN_APPROVAL_INPROCESS — citizen accepts T&C and forwards/pays.")
    add_bullet(doc, "Can SEND_TO_ARCHITECT via /editApplication/bpa|ocbpa/:tenantId/:applicationNo.")
    add_bullet(doc, "Pays BPA.NC_APP_FEE / SAN_FEE / OC fees via payment collect URLs.")
    add_bullet(doc, "My applications lists BPAREG + BPA. Cannot use Architect home unless MDMS role mapping matches.")

    add_heading(doc, "3.2 Architect / professional stakeholder", 2)
    add_para(
        doc,
        "After BPAREG is APPROVED, user receives a trade role from MDMS StakeholderRegistraition.TradeTypetoRoleMapping "
        "(e.g. BPA_ARCHITECT). They still use the citizen portal. home.js blocks users whose roles are not in that mapping.",
    )
    add_table(
        doc,
        ["Role code (typical)", "License type", "Notes"],
        [
            ("BPA_ARCHITECT", "Architect", "Council number mandatory on LicenseType step"),
            ("BPA_ENGINEER", "Engineer", "Same apply/inbox capabilities after approval"),
            ("BPA_BUILDER", "Builder", "Same"),
            ("BPA_STRUCTURALENGINEER", "Structural Engineer", "Same"),
            ("BPA_SUPERVISOR", "Supervisor", "Same"),
            ("BPA_TOWNPLANNER", "Town Planner", "Same"),
        ],
    )
    add_bullet(doc, "eDCR scrutiny (new + OC), BPA/OCBPA/PAP apply, architect inbox, eDCR inbox.")
    add_bullet(doc, "On BPA detail, if status is CITIZEN_APPROVAL_INPROCESS and user is BPA_ARCHITECT, action bar is hidden (citizen must act).")
    add_bullet(doc, "Session BPA_ARCHITECT_NAME stores trade type prefix from approved license for ack messages.")
    add_bullet(doc, "OwnerDetails blocks pure CITIZEN users from completing architect owner step.")

    add_heading(doc, "3.3 Employee", 2)
    add_table(
        doc,
        ["Role", "Inbox / UI", "Tasks"],
        [
            ("BPAREG_DOC_VERIFIER", "Stakeholder inbox", "Verify stakeholder documents"),
            ("BPAREG_APPROVER", "Stakeholder inbox", "Approve/reject professional license"),
            ("BPAREG_EMPLOYEE", "Stakeholder inbox search link", "BPAREG employee access"),
            ("BPA_VERIFIER", "BPA inbox", "Document/plan verification"),
            ("BPA_FIELD_INSPECTOR", "BPA inbox + InspectionReport", "Field inspection checklist & docs"),
            ("BPA_NOC_VERIFIER", "BPA inbox", "NOC verification"),
            ("BPA_APPROVER", "BPA inbox", "Final permit / OC approval"),
            ("CEMP", "BPA inbox + payments", "Counter employee payment collection"),
        ],
    )
    add_para(doc, "EmployeeCard hides Stakeholder or BPA links if showHidingLinksForStakeholder / showHidingLinksForBPA return empty.")

    add_heading(doc, "4. Complete route catalog", 1)
    add_heading(doc, "4.1 Citizen / Architect routes", 2)
    add_para(doc, "Base: http://localhost:3000/upyog-ui/citizen/obps")
    add_table(
        doc,
        ["Route", "Auth", "Primary role", "File / component", "Task"],
        [
            ("/home", "Private", "Architect", "pages/citizen/home.js", "Stakeholder home + apply links"),
            ("/search/application", "Private", "Architect", "pages/employee/Search.js", "Search applications"),
            ("/search/obps-application", "Private", "Architect", "pages/employee/Search.js", "Same search"),
            ("/edcrscrutiny/apply/*", "Private", "Architect", "EDCR/index.js", "New eDCR wizard"),
            ("/edcrscrutiny/apply/home", "Private", "Architect", "EDCRForm", "City, name, plan file"),
            ("/edcrscrutiny/apply/acknowledgement", "Private", "Architect", "EDCRAcknowledgement", "eDCR result"),
            ("/edcrscrutiny/oc-apply/*", "Private", "Architect", "OCEDCR/index.js", "OC eDCR wizard"),
            ("/edcrscrutiny/oc-apply/docs-required", "Private", "Architect", "OCEDCRDocsRequired", "OC info"),
            ("/edcrscrutiny/oc-apply/home", "Private", "Architect", "OCeDCRScrutiny", "Permit search"),
            ("/edcrscrutiny/oc-apply/upload-diagram", "Private", "Architect", "OCUploadPlanDiagram", "OC DXF upload"),
            ("/edcrscrutiny/oc-apply/acknowledgement", "Private", "Architect", "OCEDCRAcknowledgement", "OC eDCR result"),
            ("/bpa/:applicationType/:serviceType/*", "Private", "Architect", "NewBuildingPermit", "BPA apply wizard"),
            ("/ocbpa/:applicationType/:serviceType/*", "Private", "Architect", "OCBuildingPermit", "OC BPA wizard"),
            ("/stakeholder/apply/*", "Private", "Citizen", "StakeholderRegistration", "License apply"),
            ("/openlink/stakeholder/apply/*", "Public", "Citizen", "StakeholderRegistration", "Public license apply"),
            ("/preApprovedPlan/*", "Private", "Architect", "PreApprovedPlan", "PAP wizard"),
            ("/my-applications", "Private", "Citizen+Architect", "MyApplication", "List applications"),
            ("/bpa/inbox", "Private", "Architect", "employee/Inbox (reused)", "BPA inbox on citizen"),
            ("/edcr/inbox", "Private", "Architect", "EdcrInbox", "eDCR inbox"),
            ("/stakeholder/:id", "Private", "Citizen+Architect", "ApplicationDetail", "BPAREG detail"),
            ("/bpa/:id", "Private", "Citizen+Architect", "BpaApplicationDetail", "BPA detail + actions"),
            ("/editApplication/bpa/:tenantId/:applicationNo/*", "Private", "Citizen", "BPASendToArchitect", "Send to architect"),
            ("/editApplication/ocbpa/:tenantId/:applicationNo/*", "Private", "Citizen", "OCSendToArchitect", "OC send to architect"),
            ("/sendbacktocitizen/bpa/:tenantId/:applicationNo/*", "Private", "Architect", "BPASendBackToCitizen", "Resume after send-back"),
            ("/sendbacktocitizen/ocbpa/:tenantId/:applicationNo/*", "Private", "Architect", "OCSendBackToCitizen", "OC resume"),
            ("/response", "Private", "All", "OBPSResponse", "Post-action response / pay"),
        ],
    )

    add_heading(doc, "4.2 Employee routes", 2)
    add_para(doc, "Base: http://localhost:3000/upyog-ui/employee/obps")
    add_table(
        doc,
        ["Route", "Primary role", "File", "Task"],
        [
            ("/inbox", "BPA_* / CEMP", "Inbox/index.js", "BPA inbox"),
            ("/inbox/bpa/:id", "BPA_* / CEMP", "BpaApplicationDetails", "Process BPA"),
            ("/bpa/:id", "BPA_* / CEMP", "BpaApplicationDetails", "Same detail from other entry"),
            ("/stakeholder-inbox", "BPAREG_*", "stakeholderInbox", "License inbox"),
            ("/stakeholder-inbox/stakeholder/:id", "BPAREG_*", "ApplicationDetail", "Process license"),
            ("/search/application", "BPA or BPAREG", "Search.js", "Search"),
            ("/search/application/bpa/:id", "BPA_*", "BpaApplicationDetails", "Open from search"),
            ("/search/application/stakeholder/:id", "BPAREG_*", "ApplicationDetail", "Open from search"),
            ("/response", "Employee", "OBPSResponse", "Workflow result"),
            ("/stakeholder-response", "BPAREG_*", "StakeholderResponse", "License workflow result"),
            ("/ObpsDailyCollectionReport/*", "Employee", "EnhancedReport", "Daily collection report"),
            ("/ObpsApplicationStatusReport/*", "Employee", "EnhancedReport", "Application status report"),
        ],
    )

    add_heading(doc, "5. Wizard step maps", 1)
    add_heading(doc, "5.1 eDCR New Construction", 2)
    add_table(doc, ["Step", "Route suffix", "Component", "Functions / submit"], [
        ("1", "home", "EDCRForm", "setApplicantName, setTypeOfTenantID, selectfile, handleSubmit → CreateEDCR.handleSelect → EDCRService.create"),
        ("2", "acknowledgement", "EDCRAcknowledgement", "printReciept, routeToBPAScreen"),
    ])
    add_heading(doc, "5.2 OC eDCR", 2)
    add_table(doc, ["Step", "Route suffix", "Component", "Key functions"], [
        ("1", "docs-required", "OCEDCRDocsRequired", "goNext"),
        ("2", "home", "OCeDCRScrutiny", "setPermitNo, setOCPermitDate, getSearchResults, routeToNextPage, handleSubmit"),
        ("3", "upload-diagram", "OCUploadPlanDiagram", "selectfile, handleSubmit → createOCEdcr"),
        ("4", "acknowledgement", "OCEDCRAcknowledgement", "printReciept"),
    ])
    add_heading(doc, "5.3 Building Permit (NewConfig.js — runtime)", 2)
    add_table(doc, ["#", "Route", "Component", "Task"], [
        ("1", "docs-required", "DocsRequired", "Required docs + architect name"),
        ("2", "basic-details", "BasicDetails", "eDCR number + scrutiny summary"),
        ("3", "plot-details", "PlotDetails", "Holding / registration"),
        ("4", "scrutiny-details", "ScrutinyDetails", "Blocks / sub-occupancy"),
        ("5", "search-property", "CPTSearchProperty", "Link property tax record"),
        ("6", "search-results", "CPTSearchResults", "Pick property"),
        ("—", "create-property", "CPTCreateProperty", "Optional create PT"),
        ("—", "acknowledge-create-property", "CPTAcknowledgement", "PT create ack"),
        ("7", "location", "LocationDetails", "Address + GIS"),
        ("8", "owner-details", "OwnerDetails", "Owners; may create INITIATED BPA"),
        ("9", "document-details", "DocumentDetails", "Upload docs"),
        ("10", "noc-details", "NOCDetails", "NOC uploads"),
        ("11", "check", "BPACheckPage", "Review"),
        ("12", "acknowledgement", "BPAAcknowledgement", "Submit result"),
    ])
    add_heading(doc, "5.4 OC Building Permit", 2)
    add_table(doc, ["#", "Route", "Component"], [
        ("1", "docs-required", "DocsRequired"),
        ("2", "basic-details", "OCBasicDetails"),
        ("3", "plot-details", "PlotDetails"),
        ("4", "scrutiny-details", "ScrutinyDetails"),
        ("5", "document-details", "DocumentDetails"),
        ("6", "noc-details", "NOCDetails"),
        ("7", "check", "OCBPACheckPage"),
        ("8", "acknowledgement", "OCBPAAcknowledgement"),
    ])
    add_heading(doc, "5.5 Stakeholder registration", 2)
    add_table(doc, ["#", "Route", "Component", "Task"], [
        ("1", "stakeholder-docs-required", "StakeholderDocsRequired", "Info"),
        ("2", "provide-license-type", "LicenseType", "Trade type + council no"),
        ("3", "license-details", "LicenseDetails", "Personal details"),
        ("4", "Permanent-address", "PermanentAddress", "Permanent address"),
        ("5", "correspondence-address", "CorrospondenceAddress", "Correspondence; may BPAREGCreate"),
        ("6", "stakeholder-document-details", "StakeholderDocuments", "Uploads"),
        ("7", "check", "StakeholderCheckPage", "Review"),
        ("8", "acknowledgement", "StakeholderAcknowledgement", "License application no"),
    ])
    add_heading(doc, "5.6 Pre-Approved Plan", 2)
    add_table(doc, ["#", "Route", "Component"], [
        ("1", "documents-required", "PreApprovedDocsRequired"),
        ("2", "planDetails", "BuildingPlanScrutiny"),
        ("3", "BasicDetails", "BasicDetails"),
        ("4", "plot-details", "PlotDetails"),
        ("5", "scrutiny-details", "ScrutinyDetails"),
        ("6", "location", "LocationDetails"),
        ("7", "owner-details", "OwnerDetails"),
        ("8", "document-details", "DocumentDetails"),
        ("9", "check", "BPACheckPage"),
        ("10", "acknowledgement", "BPAAcknowledgement"),
    ])

    add_heading(doc, "6. Component registry (Module.js)", 1)
    add_para(doc, "initOBPSComponents registers:")
    add_para(
        doc,
        "OBPSModule, OBPSLinks, OBPSCard, BPACitizenHomeScreen, EDCRForm, BasicDetails, BuildingPlanScrutiny, "
        "PreApprovedDocsRequired, DocsRequired, PlotDetails, ScrutinyDetails, OwnerDetails, DocumentDetails, NOCDetails, "
        "LocationDetails, GIS, OCEDCRDocsRequired, OCeDCRScrutiny, OCUploadPlanDiagram, StakeholderDocsRequired, "
        "LicenseType, LicenseDetails, CorrospondenceAddress, PermanentAddress, StakeholderDocuments, OCBasicDetails, "
        "OBPSSearchApplication, InspectionReport, BPAInbox, StakeholderInbox, StakeholderCheckPage, BPACheckPage, "
        "OCBPACheckPage, OCBPASendBackCheckPage, EDCRAcknowledgement, OCEDCRAcknowledgement, BPAAcknowledgement, "
        "OCBPAAcknowledgement, OCSendBackAcknowledgement, StakeholderAcknowledgement, ObpsCreateEDCR, ObpsCreateOCEDCR, "
        "ObpsNewBuildingPermit, ObpsOCBuildingPermit, ObpsStakeholderRegistration, ObpsPreApprovedPlan, "
        "ObpsCitizenBpaApplicationDetail, ObpsBPASendToArchitect, ObpsOCSendToArchitect, ObpsBPASendBackToCitizen, "
        "ObpsOCSendBackToCitizen, ObpsEdcrInbox, ObpsEmpApplicationDetail, ObpsEmployeeBpaApplicationDetail, "
        "EnhancedReport, ReportSearchApplication.",
        size=10,
    )

    add_heading(doc, "7. Utils function reference (src/utils/index.js)", 1)
    add_table(
        doc,
        ["Function", "What it does"],
        [(k, v) for k, v in UTILS_FUNCS.items()],
    )

    add_heading(doc, "8. Every source file — roles, tasks, functions", 1)
    add_para(
        doc,
        "The following is a complete inventory of every .js/.jsx file in the package (excluding node_modules and generator scripts). "
        "Functions are extracted from source (function declarations and const arrows). Inner helpers listed with the file.",
    )

    last = None
    for f in files:
        parts = f["rel"].split("/")
        group = "/".join(parts[:-1]) if len(parts) > 1 else "(package root)"
        if group != last:
            add_heading(doc, f"Folder: {group}", 2)
            last = group
        catalog = TASKS.get(f["rel"], ("See path", "Implements UI/logic for this module path. See functions below."))
        role, task = catalog if isinstance(catalog, tuple) else ("All", catalog)
        role = infer_role(f["rel"], role)
        add_heading(doc, f["rel"], 3)
        add_para(doc, f"Lines: {f['lines']}    |    Role: {role}", bold=True, size=10)
        add_para(doc, f"Task: {task}")
        if f["fns"]:
            rows = []
            for name, args in f["fns"]:
                extra = ""
                if f["rel"] == "src/utils/index.js" and name in UTILS_FUNCS:
                    extra = UTILS_FUNCS[name]
                elif name in ("handleSelect", "goNext", "createApplication", "handleSkip", "onSuccess"):
                    extra = {
                        "handleSelect": "Merge step data into session params and call goNext.",
                        "goNext": "Navigate to config nextStep or /check if nextStep is null.",
                        "createApplication": "Navigate to acknowledgement (mutation runs on ack screen).",
                        "handleSkip": "No-op skip handler required by FormStep.",
                        "onSuccess": "Invalidate queries; some wizards clear session params.",
                    }[name]
                rows.append((name, args or "—", extra))
            add_table(doc, ["Function", "Parameters", "Detail"], rows)
        else:
            add_para(doc, "No JS functions (config/data export only).", size=10)
        if f["apis"]:
            add_para(doc, "API / services: " + ", ".join(f["apis"]), size=10)
        if f["hooks"]:
            add_para(doc, "Digit.Hooks: " + ", ".join(f["hooks"]), size=10)

    add_heading(doc, "9. Session storage and flags", 1)
    add_table(
        doc,
        ["Key", "Set by", "Purpose"],
        [
            ("OBPS_TENANTS", "OBPSModule", "Tenant list for module"),
            ("EDCR_CREATE", "CreateEDCR", "eDCR create result / error"),
            ("OC_EDCR_CREATE", "CreateOCEDCR", "OC eDCR result"),
            ("EDCR_BACK", "CreateEDCR / home", "IS_EDCR_BACK for inbox return"),
            ("EDCR.INBOX", "EdcrInbox reducer", "eDCR inbox search/filter/table"),
            ("BUILDING_PERMIT", "BPA/OC/stakeholder wizards", "Wizard form params"),
            ("BPA_PRE_APPROVED_CREATE", "PreApprovedPlan", "PAP wizard params"),
            ("BPA_HOME_CREATE", "home.js", "Cleared on home mount"),
            ("OBPS.INBOX", "Employee/Architect BPA inbox", "Inbox forms"),
            ("STAKEHOLDER.INBOX", "stakeholder inbox", "BPAREG inbox forms"),
            ("OBPS_PT", "NewBuildingPermit", "Property tax integration flag"),
            ("isEDCRDisable", "Ack / detail", "Lock eDCR field on basic details"),
            ("isPermitApplication", "home / ack", "Clear BUILDING_PERMIT when starting new"),
            ("clickOnBPAApplyAfterEDCR", "EDCR ack CTA", "Hide back button on docs-required"),
            ("BPA_ARCHITECT_NAME", "DocsRequired", "Trade type for messages"),
            ("BPA_SUBMIT_APP", "send-to-architect / send-back", "Ack message variant"),
            ("BPA_IS_ALREADY_WENT_OFF_DETAILS", "edit flows", "Avoid re-hydration"),
            ("BPAintermediateValue", "getBPAFormData", "Resume draft into wizard"),
            ("BPAREGintermediateValue", "LicenseType", "Resume BPAREG draft"),
            ("isEDCRAPIType", "BasicDetails", "Which scrutiny API was used"),
            ("updateData", "workflow", "Read by OBPSResponse"),
            ("EMPLOYEE_MUTATION_*", "employee BPA detail", "Mutation success tracking"),
        ],
    )

    add_heading(doc, "10. APIs and MDMS", 1)
    add_heading(doc, "10.1 HTTP / Digit services used in this package", 2)
    add_table(
        doc,
        ["Call", "Where", "Purpose"],
        [
            ("Digit.EDCRService.create", "EDCR / OCEDCR wizards", "POST /edcr/rest/dcr/scrutinize multipart"),
            ("Digit.OBPSService.scrutinyDetails", "utils, OCeDCRScrutiny, getBPAFormData", "Fetch eDCR by number"),
            ("Digit.OBPSService.BPASearch", "utils, OCeDCR, OBPSResponse", "Search BPA by edcr/app/approval no"),
            ("Digit.OBPSService.create", "OwnerDetails, ScrutinyDetails", "Create INITIATED BPA"),
            ("Digit.OBPSService.BPAREGCreate", "CorrospondenceAddress", "Draft stakeholder license"),
            ("Digit.OBPSService.edcr_report_download", "BPA details", "Permit/OC eDCR PDF"),
            ("Digit.PaymentService.fetchBill", "OBPSResponse", "Outstanding fee"),
            ("Digit.PaymentService.generatePdf / printReciept", "details, response, utils", "Receipts, revocation, permit"),
            ("Digit.UserService.userSearch / userCreate", "OwnerDetails", "Owner account"),
            ("Digit.UserService.getUser", "many", "Logged-in user + roles"),
            ("Digit.WorkflowService.getAllApplication", "OCeDCRScrutiny", "WF lookup"),
            ("PreApprovedPlanService.search", "scrutinyDetailsData", "PAP drawing (8-char no)"),
        ],
    )
    add_heading(doc, "10.2 MDMS masters", 2)
    add_table(
        doc,
        ["Module.Master", "Used for"],
        [
            ("tenant.citymodule (BPAAPPLY)", "eDCR city dropdown"),
            ("StakeholderRegistraition.TradeTypetoRoleMapping", "Home gate + LicenseType list + roles"),
            ("BPA.homePageUrlLinks", "Home BPA/OC cards + post-eDCR apply URL"),
            ("BPA.RiskTypeComputation", "LOW/MEDIUM/HIGH from plot area/floors"),
            ("BPA.DocumentTypes / DocTypeMapping", "Required and uploadable docs"),
            ("BPA.EdcrConfig / BuildingPermitConfig / OCEdcrConfig / OCBuildingPermitConfig / StakeholderConfig", "Optional step overrides"),
            ("common-masters.DocumentType", "Common document metadata"),
            ("BPA_FORM_CONFIG (getFormConfig)", "Loads all form configs from MDMS"),
        ],
    )

    add_heading(doc, "11. Workflow / business services", 1)
    add_table(
        doc,
        ["businessService", "Meaning", "Fee mapping (getBusinessServices)"],
        [
            ("BPA_LOW", "Low-risk permit", "BPA.LOW_RISK_PERMIT_FEE"),
            ("BPA", "New construction permit", "PENDING_APPL_FEE → BPA.NC_APP_FEE else BPA.NC_SAN_FEE"),
            ("BPA-PAP", "Pre-approved plan", "Same as BPA (NC_APP / NC_SAN); workflow action APPLY"),
            ("BPA_OC", "Occupancy certificate", "PENDING_APPL_FEE → BPA.NC_OC_APP_FEE else BPA.NC_OC_SAN_FEE"),
            ("BPAREG", "Stakeholder registration", "Trade license fees (BPAREG module)"),
        ],
    )
    add_para(doc, "Notable statuses: INITIATED, INPROGRESS, CITIZEN_APPROVAL_INPROCESS, PENDING_APPL_FEE, APPROVED, REJECTED, PERMIT REVOCATION, SEND_BACK_TO_CITIZEN.")
    add_para(doc, "Notable workflow actions: SEND_TO_CITIZEN, SEND_TO_ARCHITECT, APPLY (PAP), FORWARD, APPROVE, plus employee inbox actions from workflow config.")

    add_heading(doc, "12. Shared wizard function pattern", 1)
    add_para(doc, "Almost every apply wizard (EDCR, BPA, OC, Stakeholder, PAP, SendToArchitect, SendBack) implements the same functions:")
    add_bullet(doc, "handleSelect(key, data, skipStep, isFromCreateApi) — merge into useSessionStorage params.")
    add_bullet(doc, "goNext(skipStep) — read nextStep from config for current path; null → /check.")
    add_bullet(doc, "createApplication() — navigate to acknowledgement.")
    add_bullet(doc, "handleSkip() — empty.")
    add_bullet(doc, "onSuccess() — react-query invalidate.")
    add_para(doc, "Acknowledgement screens typically run the create/update mutation (useObpsAPI / BPAREG) using convertToBPAObject or convertToStakeholderObject.")

    add_heading(doc, "13. File count summary", 1)
    by_role = {"Citizen": 0, "Architect": 0, "Employee": 0, "Mixed/All": 0}
    for f in files:
        cat = TASKS.get(f["rel"])
        role = cat[0] if cat else infer_role(f["rel"], None)
        if "Employee" in role and "Citizen" not in role and "Architect" not in role:
            by_role["Employee"] += 1
        elif "Citizen" in role and "Architect" not in role and "Employee" not in role:
            by_role["Citizen"] += 1
        elif role.startswith("Architect"):
            by_role["Architect"] += 1
        else:
            by_role["Mixed/All"] += 1
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ("Total JS/JSX files", str(len(files))),
            ("Primarily Citizen", str(by_role["Citizen"])),
            ("Primarily Architect", str(by_role["Architect"])),
            ("Primarily Employee", str(by_role["Employee"])),
            ("Shared / mixed / all", str(by_role["Mixed/All"])),
            ("Total extracted functions", str(sum(len(f["fns"]) for f in files))),
        ],
    )

    foot = doc.add_paragraph()
    r = foot.add_run(
        "Generated from live source under packages/modules/obps. Regenerate with generate_obps_complete_docx.py. "
        "Related: OBPS_eDCR_Scrutiny_High_Level_Flow.docx, OBPS_eDCR_Scrutiny_Low_Level_Diagram.docx."
    )
    set_run_font(r, 9, False, (120, 120, 120))

    doc.save(OUT)
    print(f"Wrote {OUT} files={len(files)} functions={sum(len(f['fns']) for f in files)}")


if __name__ == "__main__":
    main()
